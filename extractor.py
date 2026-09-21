"""Utilities for extracting NBTC EMF reports into a consistent table."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
import re
import unicodedata
from typing import Iterable, Sequence

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

try:
    from PIL import ImageOps
    import pypdfium2 as pdfium
    import pytesseract
except ImportError:
    ImageOps = None
    pdfium = None
    pytesseract = None


REPORT_COLUMNS = [
    "ลำดับที่",
    "ผู้ประกอบการ",
    "เลขที่ใบอนุญาตตั้ง",
    "ที่ตั้ง",
    "ตำบล",
    "อำเภอ",
    "จังหวัด",
    "รหัสไปรษณีย์",
    "Longitude",
    "Latitude",
    "ความถี่",
    "ตราอักษร",
    "รุ่น/แบบ",
    "กำลังส่ง (วัตต์)",
    "อัตราขยายสายอากาศ (dBi)",
    "ความสูงสายอากาศ (เมตร)",
    "ระยะห่างจากเสา ที่ตั้งสายอากาศ",
    "ระยะที่วัด/คำนวณ (เมตร)",
    "ระดับการแผ่คลื่นแม่เหล็กไฟฟ้าสูงสุด",
    "วันที่วัด/คำนวณ",
    "ลงชื่อ",
    "วันที่รายงาน",
]

QA_COLUMNS = [
    "ไฟล์ต้นทาง",
    "หน้า",
    "วิธีอ่าน PDF",
    "สถานะ",
    "ข้อมูลที่ไม่พบ",
    "จำนวนความถี่",
    "หมายเหตุ",
]

PAGE_DATA_COLUMNS = [
    "ไฟล์ต้นทาง",
    "หน้าต้นทาง",
    "หน้าที่พบ",
    "ผู้ประกอบการ",
    "ที่ตั้ง",
    "พิกัด",
    "ค่าแผ่คลื่นสูงสุด",
    "ชื่อ",
    "ตำแหน่ง",
    "วันที่",
    "วันที่ประกาศ",
    "ชนิดโครงสร้างเสาส่งสัญญาณโทรคมนาคม",
    "ลำดับ",
    "หมู่",
    "สถานะ",
    "ข้อมูลที่ไม่พบ",
]

MAIN_REPORT_PAGE_TYPE = "แบบรายงานระดับการแผ่คลื่นแม่เหล็กไฟฟ้า ของสถานีวิทยุคมนาคม"
PAGE_TYPE_REQUIREMENTS = {
    MAIN_REPORT_PAGE_TYPE: ("ผู้ประกอบการ", "ที่ตั้ง", "พิกัด", "ค่าแผ่คลื่นสูงสุด", "ชื่อ", "วันที่"),
    "หลักฐานการทำความเข้าใจ": ("ที่ตั้ง", "วันที่ประกาศ"),
    "ข้อมูลทำความเข้าใจเกี่ยวกับการติดตั้งสถานีวิทยุ-ส่งสัญญาณโทรศัพท์เคลื่อนที่": (
        "ที่ตั้ง",
        "ชนิดโครงสร้างเสาส่งสัญญาณโทรคมนาคม",
        "วันที่",
    ),
    "รูปถ่ายการแจกเอกสาร": ("ชื่อ", "ที่ตั้ง"),
    "บันทึกการประชุม": ("ชื่อ", "ที่ตั้ง", "ตำแหน่ง"),
    "เอกสารลงทะเบียนประชุม": ("ลำดับ", "หมู่"),
    "ภาพถ่ายการประชุม": ("วันที่",),
    "เขียนที่": ("ที่ตั้ง", "ชื่อ"),
}

OCR_DPI = 300
OCR_LANGUAGE = "tha+eng"


# Some Thai PDF fonts omit their ToUnicode mapping.  pdfplumber then exposes
# the missing marks as ``(cid:...)`` while pypdf exposes Latin characters.
# These four mappings are the Thai tone marks used by the supplied report.
_THAI_FONT_MARK_FIXES = {
    "(cid:201)": "่",
    "(cid:202)": "้",
    "(cid:203)": "๊",
    "(cid:205)": "์",
    "É": "่",
    "Ê": "้",
    "Ë": "๊",
    "Í": "์",
}

# pdfplumber can preserve the visual spacing found in some Thai PDF fonts.
# These replacements are deliberately limited to terms used as field labels.
_THAI_TYPO_FIXES = {
    "ต าบล": "ตำบล",
    "ตา บล": "ตำบล",
    "อ าเภอ": "อำเภอ",
    "อา เภอ": "อำเภอ",
    "จา กดั": "จำกัด",
    "ที่ต้งั": "ที่ตั้ง",
    "ที่ ตั้ ง": "ที่ตั้ง",
    "หนา้ สา รวจ": "หน้าสำรวจ",
    "วนั ที่": "วันที่",
    "วันที่วดั": "วันที่วัด",
    "คา นวณ": "คำนวณ",
    "ค ำนวณ": "คำนวณ",
    "คำ นวณ": "คำนวณ",
    "ผมู้ ีอา นาจ": "ผู้มีอำนาจ",
    "ผู้มีอ ำนาจ": "ผู้มีอำนาจ",
    "ผมู้ ีอำ นำจ": "ผู้มีอำนาจ",
    "กระทา การ": "กระทำการ",
    "บริษทั": "บริษัท",
    "รหัสไปรษณยี ์": "รหัสไปรษณีย์",
}


LABELS = {
    "operator": ("หน่วยงาน",),
    "license": ("เลขที่ใบอนุญาตตั้ง", "เลขที่ใบอนุญาตต้งั"),
    "location": ("ที่ตั้ง", "ที่ต้งั"),
    "subdistrict": ("ตำบล", "ตาบล"),
    "district": ("อำเภอ", "อาเภอ"),
    "province": ("จังหวัด",),
    "zipcode": ("รหัสไปรษณีย์", "รหัสไปรษณยี"),
    "longitude": ("longitude", "longtitude"),
    "latitude": ("latitude",),
    "date_calc": ("วันที่วัด/คำนวณ", "วันที่วัด/คานวณ"),
    "signature": ("ผู้มีอำนาจลงนาม", "ผู้มีอำนาจลงนำม"),
    "date_report": ("วันที่รายงาน",),
}

MISSING_FIELD_NAMES = {
    "operator": "ผู้ประกอบการ",
    "license": "เลขที่ใบอนุญาตตั้ง",
    "location": "ที่ตั้ง",
    "subdistrict": "ตำบล",
    "district": "อำเภอ",
    "province": "จังหวัด",
    "zipcode": "รหัสไปรษณีย์",
    "longitude": "Longitude",
    "latitude": "Latitude",
    "frequencies": "ความถี่",
    "max_dist_text": "ระยะห่างจากเสา ที่ตั้งสายอากาศ",
    "max_dist_val": "ระยะที่วัด/คำนวณ",
    "max_rad": "ระดับการแผ่คลื่นแม่เหล็กไฟฟ้าสูงสุด",
    "date_calc": "วันที่วัด/คำนวณ",
    "signature": "ลงชื่อ",
    "date_report": "วันที่รายงาน",
}

FREQUENCY_FIELD_NAMES = {
    "brand": "ตราอักษร",
    "model": "รุ่น/แบบ",
    "power": "กำลังส่ง (วัตต์)",
    "gain": "อัตราขยายสายอากาศ (dBi)",
    "height": "ความสูงสายอากาศ (เมตร)",
}


@dataclass
class StationPage:
    """A parsed report page and the review information associated with it."""

    source_file: str
    page_number: int
    data: dict[str, str]
    frequencies: list[dict[str, str]]
    missing_fields: list[str]
    parser_method: str = "pdfplumber"
    note: str = ""


@dataclass
class PageDataRecord:
    """One data record requested for a recognised page type in the document."""

    source_file: str
    page_number: int
    page_type: str
    data: dict[str, str]
    required_fields: Sequence[str]
    parser_method: str = "pypdf"
    note: str = ""


def clean_text(value: object) -> str:
    """Return a readable, single-line value without changing factual content."""
    if value is None:
        return ""

    text = str(value)
    text = text.replace("\u00a0", " ").replace("\n", " ")
    for incorrect, corrected in _THAI_FONT_MARK_FIXES.items():
        text = text.replace(incorrect, corrected)
    # Thai Sara Am may be emitted as either of its decomposed visual orders.
    text = re.sub(r"\u0e4d\s*\u0e32|\u0e32\s*\u0e4d", "\u0e33", text)
    text = unicodedata.normalize("NFC", text)
    for incorrect, corrected in _THAI_TYPO_FIXES.items():
        text = text.replace(incorrect, corrected)
    return re.sub(r"\s+", " ", text).strip()


def label_key(value: object) -> str:
    """Normalise labels only, so harmless font-spacing differences still match."""
    text = clean_text(value).lower().replace("longtitude", "longitude")
    text = re.sub(r"[\s:()\[\]{}=_\-/]", "", text)
    # Thai combining marks are frequently emitted in a different order by PDFs.
    text = re.sub(r"[\u0e31-\u0e3a\u0e47-\u0e4e]", "", text)
    return text


def matches_label(value: object, aliases: Sequence[str]) -> bool:
    key = label_key(value)
    return bool(key) and any(label_key(alias) in key for alias in aliases)


def is_same_label(value: object, aliases: Sequence[str]) -> bool:
    key = label_key(value)
    return bool(key) and any(key == label_key(alias) for alias in aliases)


def _contains_label(text: object, aliases: Sequence[str]) -> bool:
    """Check whether a block of text contains one of the supplied headings."""
    key = label_key(text)
    return bool(key) and any(label_key(alias) in key for alias in aliases)


def _normalise_license(value: str) -> str:
    """Keep only genuine licence values, never the next empty-table label."""
    value = clean_text(value)
    if _contains_label(value, (*LABELS["location"], *LABELS["subdistrict"], *LABELS["district"])):
        return ""
    return value if re.search(r"\d{4,}", value) else ""


def clean_rows(tables: Iterable[Sequence[Sequence[object]]]) -> list[list[str]]:
    rows: list[list[str]] = []
    for table in tables:
        for row in table or []:
            cleaned = [clean_text(cell) for cell in (row or [])]
            if any(cleaned):
                rows.append(cleaned)
    return rows


def value_after_label(rows: Sequence[Sequence[str]], aliases: Sequence[str]) -> str:
    """Read a label's value from the same row, with a one-row fallback."""
    for row_index, row in enumerate(rows):
        for column_index, cell in enumerate(row):
            if not matches_label(cell, aliases):
                continue

            for candidate in row[column_index + 1 :]:
                if candidate and not is_same_label(candidate, aliases):
                    return candidate

            # Some report variants put the value in the row directly below.
            if row_index + 1 < len(rows):
                next_row = rows[row_index + 1]
                candidates = next_row[column_index:] + next_row[:column_index]
                for candidate in candidates:
                    if candidate and not is_same_label(candidate, aliases):
                        return candidate
    return ""


def is_frequency(value: str) -> bool:
    return bool(re.fullmatch(r"\d{2,5}(?:\.\d+)?", value.strip()))


def extract_frequency_rows(rows: Sequence[Sequence[str]]) -> list[dict[str, str]]:
    """Extract the six frequency-table fields from either supplied table layout."""
    frequencies: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str, str, str]] = set()

    for row in rows:
        populated = [cell for cell in row if cell]
        if len(populated) < 6 or not is_frequency(populated[0]):
            continue

        values = tuple(populated[:6])
        if values in seen:
            continue
        seen.add(values)
        frequencies.append(
            {
                "frequency": values[0],
                "brand": values[1],
                "model": values[2],
                "power": values[3],
                "gain": values[4],
                "height": values[5],
            }
        )
    return frequencies


def is_numeric_value(value: str) -> bool:
    return bool(re.fullmatch(r"\d+(?:\.\d+)?%?", value.strip()))


def extract_maximum(rows: Sequence[Sequence[str]]) -> tuple[str, str, str]:
    """Find the final summary row rather than the similarly named table heading."""
    label = "ระดับการแผ่คลื่นแม่เหล็กไฟฟ้าสูงสุด"
    candidates: list[tuple[str, str]] = []

    for row in rows:
        has_summary_label = any(label_key(label) in label_key(cell) for cell in row)
        numeric_values = [cell for cell in row if is_numeric_value(cell)]
        if has_summary_label and len(numeric_values) >= 2:
            candidates.append((numeric_values[0], numeric_values[1]))

    if candidates:
        distance, radiation = candidates[-1]
        return "ระดับสูงสุด", distance, radiation
    return "", "", ""


def join_frequency_values(frequencies: Sequence[dict[str, str]], field: str) -> str:
    values = [frequency[field] for frequency in frequencies if frequency.get(field)]
    return ", ".join(dict.fromkeys(values))


def find_missing_fields(data: dict[str, str], frequencies: Sequence[dict[str, str]]) -> list[str]:
    """List every required value that was not extracted, including partial rows."""
    checks = {**data, "frequencies": "yes" if frequencies else ""}
    missing = [
        display_name
        for field, display_name in MISSING_FIELD_NAMES.items()
        if not checks.get(field, "")
    ]

    if frequencies:
        incomplete_columns = [
            display_name
            for field, display_name in FREQUENCY_FIELD_NAMES.items()
            if any(not frequency.get(field, "") for frequency in frequencies)
        ]
        if incomplete_columns:
            missing.append(f"รายละเอียดความถี่: {', '.join(incomplete_columns)}")
    return missing


def parse_page(page: pdfplumber.page.Page, source_file: str, page_number: int) -> StationPage:
    rows = clean_rows(page.extract_tables())
    data = {field: value_after_label(rows, aliases) for field, aliases in LABELS.items()}
    data["license"] = _normalise_license(data["license"])
    max_text, max_distance, max_radiation = extract_maximum(rows)
    data.update(
        {
            "max_dist_text": max_text,
            "max_dist_val": max_distance,
            "max_rad": max_radiation,
        }
    )
    frequencies = extract_frequency_rows(rows)

    note = ""
    if not rows:
        note = "ไม่พบตารางข้อมูลที่อ่านได้ในหน้านี้"

    missing_fields = find_missing_fields(data, frequencies)
    return StationPage(
        source_file,
        page_number,
        data,
        frequencies,
        missing_fields,
        parser_method="pdfplumber",
        note=note,
    )


def _pypdf_layout_lines(page: object) -> list[str]:
    """Read one page with pypdf's layout extractor, with a version fallback."""
    try:
        text = page.extract_text(extraction_mode="layout") or ""
    except TypeError:
        text = page.extract_text() or ""
    return [clean_text(line) for line in text.splitlines() if clean_text(line)]


def _pypdf_value_after_label(lines: Sequence[str], aliases: Sequence[str]) -> str:
    """Find a field value on the same visual line, then try the next line."""
    for line_number, line in enumerate(lines):
        compact_line = re.sub(r"\s*/\s*", "/", line)
        for alias in aliases:
            compact_alias = re.sub(r"\s*/\s*", "/", clean_text(alias))
            position = compact_line.find(compact_alias)
            if position >= 0:
                value = compact_line[position + len(compact_alias) :].strip(" :")
                if value:
                    return value

        # A colon is useful for labels that pypdf split into visual glyph runs.
        if matches_label(line, aliases) and ":" in line:
            value = line.split(":", maxsplit=1)[1].strip()
            if value:
                return value
        if matches_label(line, aliases) and line_number + 1 < len(lines):
            return lines[line_number + 1].strip(" :")
    return ""


def _trim_before_next_label(value: str, following_labels: Sequence[str]) -> str:
    """Keep the value before a following field label on the same visual line."""
    positions = [value.find(label) for label in following_labels if value.find(label) >= 0]
    return value[: min(positions)].strip(" :") if positions else value.strip(" :")


def _extract_pypdf_coordinates(lines: Sequence[str]) -> tuple[str, str]:
    """Extract longitude and latitude together, as the layout reader keeps them joined."""
    coordinate_patterns = (
        r"\d{1,3}(?:\.\d+)?\s*°\s*\d{1,2}'\s*\d{1,2}\"",
        r"\b\d{1,3}\s+\d{1,2}\s+\d{1,2}\b",
    )
    for line in lines:
        if "longitude" not in line.lower() and "longtitude" not in line.lower():
            continue
        coordinates: list[str] = []
        for pattern in coordinate_patterns:
            coordinates = re.findall(pattern, line)
            if len(coordinates) >= 2:
                break
        if len(coordinates) >= 2:
            return coordinates[0], coordinates[1]
    return "", ""


def _extract_pypdf_frequency_rows(lines: Sequence[str]) -> list[dict[str, str]]:
    """Recover equipment details when pypdf compresses visual table columns."""
    frequencies: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()

    for line_number, line in enumerate(lines):
        # A true equipment row starts with the frequency.  This prevents the
        # distance-summary rows (5 m to 50 m, etc.) from becoming frequencies.
        match = re.match(
            r"^\s*(\d{2,5}(?:\.\d+)?)\s+(\d+(?:\.\d+)?)\s+"
            r"(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)(.*)$",
            line,
        )
        if not match or not is_frequency(match.group(1)):
            continue
        frequency, power, gain, height, equipment = match.groups()
        brand = ""
        model = ""

        # pypdf can join the final numeric column to the next two columns,
        # for example: "9.5ERICSSONRadio 4432 B28 (KRC 161".
        equipment_match = re.match(r"([A-Z][A-Z0-9 .&-]*?)([A-Z][a-z].*)$", equipment)
        if equipment_match:
            brand = equipment_match.group(1).strip()
            model = equipment_match.group(2).strip()
            if line_number + 1 < len(lines):
                continuation = lines[line_number + 1]
                if continuation and not continuation.startswith("ผลรวม"):
                    model = f"{model} {continuation}".strip()

        key = (frequency, power, gain, height)
        if key in seen:
            continue
        seen.add(key)
        frequencies.append(
            {
                "frequency": frequency,
                "brand": brand,
                "model": model,
                "power": power,
                "gain": gain,
                "height": height,
            }
        )
    return frequencies


def _clean_pypdf_location(value: str) -> str:
    """Remove the empty หมู่ที่ cell that pypdf places before a location value."""
    if value.startswith("หมู่ที่"):
        return value.removeprefix("หมู่ที่").strip(" :")
    return value


def _extract_pypdf_maximum(lines: Sequence[str]) -> tuple[str, str, str]:
    label = "ระดับการแผ่คลื่นแม่เหล็กไฟฟ้าสูงสุด"
    candidates: list[tuple[str, str]] = []
    for line in lines:
        if not matches_label(line, (label,)):
            continue
        numbers = re.findall(r"\d+(?:\.\d+)?%?", line)
        if len(numbers) >= 2:
            candidates.append((numbers[-2], numbers[-1]))
    if candidates:
        distance, radiation = candidates[-1]
        return "ระดับสูงสุด", distance, radiation
    return "", "", ""


def parse_text_lines(
    lines: Sequence[str], source_file: str, page_number: int, parser_method: str = "pypdf"
) -> StationPage:
    """Extract a report page from text lines supplied by a PDF reader or OCR."""
    data = {
        field: _pypdf_value_after_label(lines, aliases)
        for field, aliases in LABELS.items()
    }
    data["license"] = _normalise_license(data["license"])
    data["location"] = _clean_pypdf_location(data["location"])
    data["subdistrict"] = _trim_before_next_label(
        data["subdistrict"], ("อำเภอ", "จังหวัด", "รหัสไปรษณีย์")
    )
    data["district"] = _trim_before_next_label(
        data["district"], ("จังหวัด", "รหัสไปรษณีย์")
    )
    data["province"] = _trim_before_next_label(data["province"], ("รหัสไปรษณีย์",))
    longitude, latitude = _extract_pypdf_coordinates(lines)
    data["longitude"] = longitude
    data["latitude"] = latitude
    max_text, max_distance, max_radiation = _extract_pypdf_maximum(lines)
    data.update(
        {
            "max_dist_text": max_text,
            "max_dist_val": max_distance,
            "max_rad": max_radiation,
        }
    )
    frequencies = _extract_pypdf_frequency_rows(lines)
    note = ""
    if not lines:
        note = "ไม่พบข้อความที่อ่านได้ในหน้านี้"
    elif frequencies and any(not row["brand"] or not row["model"] for row in frequencies):
        note = "อ่านตารางความถี่แบบข้อความ: โปรดตรวจสอบตราอักษรและรุ่น/แบบ"

    return StationPage(
        source_file,
        page_number,
        data,
        frequencies,
        find_missing_fields(data, frequencies),
        parser_method=parser_method,
        note=note,
    )


def parse_pypdf_page(page: object, source_file: str, page_number: int) -> StationPage:
    """Extract a report page as text for PDFs whose table font map is incomplete."""
    return parse_text_lines(_pypdf_layout_lines(page), source_file, page_number, "pypdf")


def _ocr_page_lines(source_bytes: bytes, page_number: int) -> list[str]:
    """Render one PDF page and read Thai and English text with Tesseract OCR."""
    if not all((pytesseract, pdfium, ImageOps)):
        raise RuntimeError(
            "ไม่พบไลบรารี OCR: ติดตั้ง pytesseract, pypdfium2 และ Pillow แล้ว redeploy"
        )

    document = page = bitmap = None
    try:
        document = pdfium.PdfDocument(source_bytes)
        page = document[page_number - 1]
        bitmap = page.render(scale=OCR_DPI / 72)
        image = ImageOps.autocontrast(ImageOps.grayscale(bitmap.to_pil()))
        text = pytesseract.image_to_string(
            image,
            lang=OCR_LANGUAGE,
            config="--oem 1 --psm 6",
        )
    except pytesseract.TesseractNotFoundError as error:
        raise RuntimeError(
            "ไม่พบโปรแกรม Tesseract OCR: ติดตั้ง tesseract-ocr และชุดภาษาไทย (tha)"
        ) from error
    except pytesseract.TesseractError as error:
        raise RuntimeError(
            "OCR อ่านภาษาไทยไม่ได้: ตรวจการติดตั้งชุดภาษา tha และ eng ของ Tesseract"
        ) from error
    finally:
        for resource in (bitmap, page, document):
            close = getattr(resource, "close", None)
            if close:
                close()

    return [clean_text(line) for line in text.splitlines() if clean_text(line)]


def _append_records_from_text_lines(
    stations: list[StationPage],
    page_data_records: list[PageDataRecord],
    lines: Sequence[str],
    source_file: str,
    page_number: int,
    parser_method: str,
) -> bool:
    """Append the requested records for a recognised page and report whether it matched."""
    if _page_type(lines) != MAIN_REPORT_PAGE_TYPE:
        page_data_records.extend(
            extract_page_data_records(lines, source_file, page_number, parser_method=parser_method)
        )
        return bool(_page_type(lines))

    station = parse_text_lines(lines, source_file, page_number, parser_method)
    if any(station.data.values()) or station.frequencies:
        stations.append(station)
        page_data_records.extend(
            extract_page_data_records(lines, source_file, page_number, station, station.parser_method)
        )
        return True
    return False


def _parse_ocr_document(
    source_file: str, source_bytes: bytes
) -> tuple[list[StationPage], list[PageDataRecord], list[str]]:
    """Use OCR on every page when the user explicitly selects OCR mode."""
    stations: list[StationPage] = []
    page_data_records: list[PageDataRecord] = []
    errors: list[str] = []
    try:
        reader = PdfReader(BytesIO(source_bytes))
        for page_number in range(1, len(reader.pages) + 1):
            try:
                lines = _ocr_page_lines(source_bytes, page_number)
                if not _append_records_from_text_lines(
                    stations,
                    page_data_records,
                    lines,
                    source_file,
                    page_number,
                    "OCR (Tesseract)",
                ) and lines:
                    errors.append(
                        f"{source_file} หน้า {page_number}: OCR อ่านได้ แต่ไม่พบหัวข้อหน้าที่กำหนด"
                    )
            except Exception as error:
                errors.append(f"{source_file} หน้า {page_number}: OCR ไม่สำเร็จ - {error}")
    except Exception as error:
        errors.append(f"เปิดไฟล์ {source_file} ไม่สำเร็จ: {error}")
    return stations, page_data_records, errors


def _should_try_text_fallback(station: StationPage) -> bool:
    """Use pypdf only where the table reader has something important missing."""
    return any(
        not station.data.get(field, "")
        for field in (
            "license",
            "location",
            "max_dist_text",
            "max_dist_val",
            "max_rad",
            "date_calc",
            "date_report",
        )
    )


def merge_station_pages(table_station: StationPage, text_station: StationPage) -> StationPage:
    """Fill only blank table-reader fields; never replace already extracted values."""
    data = table_station.data.copy()
    recovered_fields: list[str] = []
    for field, value in text_station.data.items():
        if not data.get(field, "") and value:
            data[field] = value
            recovered_fields.append(field)

    frequencies = table_station.frequencies or text_station.frequencies
    notes = [note for note in (table_station.note, text_station.note) if note]
    if recovered_fields:
        display_names = {
            key: MISSING_FIELD_NAMES.get(key, key) for key in recovered_fields
        }
        notes.append(f"ใช้ pypdf เติมข้อมูล: {', '.join(display_names[key] for key in recovered_fields)}")

    return StationPage(
        table_station.source_file,
        table_station.page_number,
        data,
        frequencies,
        find_missing_fields(data, frequencies),
        parser_method="pdfplumber + pypdf" if recovered_fields else "pdfplumber",
        note=" | ".join(dict.fromkeys(notes)),
    )


_THAI_MONTH_PATTERN = (
    r"(?:มกราคม|กุมภาพันธ์|มีนาคม|เมษายน|พฤษภาคม|มิถุนายน|กรกฎาคม|"
    r"สิงหาคม|กันยายน|ตุลาคม|พฤศจิกายน|ธันวาคม)"
)
_THAI_DATE_PATTERN = re.compile(rf"\b\d{{1,2}}\s+{_THAI_MONTH_PATTERN}\s+25\d{{2}}\b")
_PERSON_PREFIX = r"(?:นาย|นางสาว|นาง|ด\.ช\.|ด\.ญ\.)"
_PERSON_PATTERN = re.compile(rf"^{_PERSON_PREFIX}\s*[ก-๙][ก-๙\s.'-]{{1,80}}$")
_PERSON_VALUE_PATTERN = re.compile(
    rf"{_PERSON_PREFIX}\s*[ก-๙][ก-๙\s.'-]*?(?=\s+{_PERSON_PREFIX}(?=[ก-๙])|$)"
)
_ADDRESS_TERMS = ("เลขที่", "หมู่ที่", "ตำบล", "อำเภอ", "จังหวัด", "ถนน", "ซอย")


def _first_thai_date(lines: Sequence[str]) -> str:
    for line in lines:
        match = _THAI_DATE_PATTERN.search(line)
        if match:
            return match.group(0)
    return ""


def _date_after_label(lines: Sequence[str], aliases: Sequence[str]) -> str:
    for index, line in enumerate(lines):
        if not matches_label(line, aliases):
            continue
        date = _first_thai_date([line, *lines[index + 1 : index + 3]])
        if date:
            return date
    return ""


def _page_location(lines: Sequence[str]) -> str:
    location = _pypdf_value_after_label(
        lines, ("สถานที่ตั้ง", "สถานที่ติดตั้ง", "ที่ตั้ง")
    )
    return _trim_before_next_label(
        location,
        ("ตำบล", "อำเภอ", "จังหวัด", "รหัสไปรษณีย์", "ชนิดโครงสร้าง"),
    )


def _person_name(lines: Sequence[str]) -> str:
    for line in lines:
        if "ชื่อสถานี" in label_key(line):
            continue
        names = _person_values(line)
        if names:
            return names[0]

    name = _pypdf_value_after_label(
        lines, ("ชื่อผู้เข้าร่วมประชุม", "ชื่อ-นามสกุล", "ชื่อ", "ผู้มีอำนาจลงนาม")
    )
    return "" if "สถานี" in name else name


def _person_values(line: str) -> list[str]:
    return [clean_text(match.group(0)) for match in _PERSON_VALUE_PATTERN.finditer(line)]


def _split_repeated_address(value: str, marker: str, expected_parts: int) -> list[str]:
    positions = [match.start() for match in re.finditer(re.escape(marker), value)]
    if len(positions) != expected_parts:
        return [value] * expected_parts
    return [value[start:end].strip() for start, end in zip(positions, [*positions[1:], len(value)])]


def _people_with_locations(lines: Sequence[str]) -> list[dict[str, str]]:
    """Return one record per named recipient in a handout-photo page."""
    people: list[dict[str, str]] = []
    for index, line in enumerate(lines):
        names = _person_values(line)
        if not names:
            continue
        address_groups: list[list[str]] = [[] for _ in names]
        for candidate in lines[index + 1 : index + 5]:
            if _person_values(candidate):
                break
            if any(term in candidate for term in _ADDRESS_TERMS):
                marker = "เลขที่" if candidate.count("เลขที่") > 1 else "ตำบล"
                parts = _split_repeated_address(candidate, marker, len(names))
                for name_index, part in enumerate(parts):
                    address_groups[name_index].append(part)
        people.extend(
            {"ชื่อ": name, "ที่ตั้ง": " ".join(address_groups[name_index])}
            for name_index, name in enumerate(names)
        )
    return people


def _registration_rows(lines: Sequence[str]) -> list[dict[str, str]]:
    """Read explicit sequence-and-village pairs where a registration table exposes them."""
    rows: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for line in lines:
        match = re.search(r"(?:^|\s)(\d{1,3})\s+.*?หมู่ที่\s*(\d{1,2})(?:\s|$)", line)
        if not match:
            continue
        values = (match.group(1), match.group(2))
        if values not in seen:
            seen.add(values)
            rows.append({"ลำดับ": values[0], "หมู่": values[1]})
    return rows


def _page_type(lines: Sequence[str]) -> str:
    text = " ".join(lines)
    if _contains_label(text, ("แบบรายงานระดับการแผ่คลื่นแม่เหล็กไฟฟ้า",)):
        return MAIN_REPORT_PAGE_TYPE
    if _contains_label(text, ("หลักฐานการทำความเข้าใจ",)):
        return "หลักฐานการทำความเข้าใจ"
    if _contains_label(text, ("รูปถ่ายการแจกเอกสาร",)):
        return "รูปถ่ายการแจกเอกสาร"
    if _contains_label(text, ("ภาพถ่ายการประชุม",)):
        return "ภาพถ่ายการประชุม"
    if _contains_label(text, ("เอกสารลงทะเบียนประชุม", "ลงทะเบียนเข้าร่วมประชุม")):
        return "เอกสารลงทะเบียนประชุม"
    if _contains_label(text, ("บันทึกการประชุม",)):
        return "บันทึกการประชุม"
    if _contains_label(text, ("เขียนที่",)):
        return "เขียนที่"
    if _contains_label(
        text,
        (
            "ข้อมูลทำความเข้าใจเกี่ยวกับการติดตั้งสถานีวิทยุ",
            "เอกสารเผยแพร่ทำความเข้าใจกับประชาชน",
            "การติดตั้งสถานีฐานโทรศัพท์เคลื่อนที่",
        ),
    ):
        return "ข้อมูลทำความเข้าใจเกี่ยวกับการติดตั้งสถานีวิทยุ-ส่งสัญญาณโทรศัพท์เคลื่อนที่"
    return ""


def _new_page_record(
    source_file: str,
    page_number: int,
    page_type: str,
    data: dict[str, str],
    parser_method: str,
    note: str = "",
) -> PageDataRecord:
    return PageDataRecord(
        source_file=source_file,
        page_number=page_number,
        page_type=page_type,
        data=data,
        required_fields=PAGE_TYPE_REQUIREMENTS[page_type],
        parser_method=parser_method,
        note=note,
    )


def extract_page_data_records(
    lines: Sequence[str],
    source_file: str,
    page_number: int,
    station: StationPage | None = None,
    parser_method: str = "pypdf",
) -> list[PageDataRecord]:
    """Extract only the user-selected fields for each recognised document page."""
    page_type = _page_type(lines)
    if not page_type:
        return []

    if page_type == MAIN_REPORT_PAGE_TYPE and station:
        data = station.data
        coordinates = ", ".join(
            value for value in (data.get("longitude", ""), data.get("latitude", "")) if value
        )
        return [
            _new_page_record(
                source_file,
                page_number,
                page_type,
                {
                    "ผู้ประกอบการ": data.get("operator", ""),
                    "ที่ตั้ง": data.get("location", ""),
                    "พิกัด": coordinates,
                    "ค่าแผ่คลื่นสูงสุด": data.get("max_rad", ""),
                    "ชื่อ": data.get("signature", ""),
                    "วันที่": data.get("date_report", "") or data.get("date_calc", ""),
                },
                station.parser_method,
                station.note,
            )
        ]

    if page_type == "หลักฐานการทำความเข้าใจ":
        return [
            _new_page_record(
                source_file,
                page_number,
                page_type,
                {
                    "ที่ตั้ง": _page_location(lines),
                    "วันที่ประกาศ": _date_after_label(lines, ("วันที่ประกาศ",)),
                },
                parser_method,
            )
        ]

    if page_type == "ข้อมูลทำความเข้าใจเกี่ยวกับการติดตั้งสถานีวิทยุ-ส่งสัญญาณโทรศัพท์เคลื่อนที่":
        return [
            _new_page_record(
                source_file,
                page_number,
                page_type,
                {
                    "ที่ตั้ง": _page_location(lines),
                    "ชนิดโครงสร้างเสาส่งสัญญาณโทรคมนาคม": _pypdf_value_after_label(
                        lines,
                        (
                            "ชนิดโครงสร้างเสาส่งสัญญาณโทรคมนาคม",
                            "ชนิดโครงสร้างเสาส่งสัญญาณ",
                            "ชนิดโครงสร้าง",
                        ),
                    ),
                    "วันที่": _first_thai_date(lines),
                },
                parser_method,
            )
        ]

    if page_type == "รูปถ่ายการแจกเอกสาร":
        people = _people_with_locations(lines)
        if people:
            return [
                _new_page_record(source_file, page_number, page_type, person, parser_method)
                for person in people
            ]
        return [_new_page_record(source_file, page_number, page_type, {}, parser_method)]

    if page_type == "บันทึกการประชุม":
        return [
            _new_page_record(
                source_file,
                page_number,
                page_type,
                {
                    "ชื่อ": _person_name(lines),
                    "ที่ตั้ง": _page_location(lines),
                    "ตำแหน่ง": _pypdf_value_after_label(lines, ("ตำแหน่ง",)),
                },
                parser_method,
            )
        ]

    if page_type == "เอกสารลงทะเบียนประชุม":
        rows = _registration_rows(lines)
        if rows:
            return [
                _new_page_record(source_file, page_number, page_type, row, parser_method)
                for row in rows
            ]
        return [
            _new_page_record(
                source_file,
                page_number,
                page_type,
                {
                    "ลำดับ": _pypdf_value_after_label(lines, ("ลำดับ",)),
                    "หมู่": _pypdf_value_after_label(lines, ("หมู่", "หมู่ที่")),
                },
                parser_method,
            )
        ]

    if page_type == "ภาพถ่ายการประชุม":
        return [
            _new_page_record(
                source_file, page_number, page_type, {"วันที่": _first_thai_date(lines)}, parser_method
            )
        ]

    return [
        _new_page_record(
            source_file,
            page_number,
            page_type,
            {"ที่ตั้ง": _page_location(lines), "ชื่อ": _person_name(lines)},
            parser_method,
        )
    ]


def parse_document_bytes(
    source_file: str, source_bytes: bytes, parser_mode: str = "auto"
) -> tuple[list[StationPage], list[PageDataRecord], list[str]]:
    """Parse selected data from main-report and recognised supporting-document pages."""
    if parser_mode not in {"auto", "pdfplumber", "pypdf", "ocr"}:
        raise ValueError("parser_mode ต้องเป็น auto, pdfplumber, pypdf หรือ ocr")

    if parser_mode == "ocr":
        return _parse_ocr_document(source_file, source_bytes)

    stations: list[StationPage] = []
    page_data_records: list[PageDataRecord] = []
    errors: list[str] = []

    if parser_mode == "pypdf":
        try:
            reader = PdfReader(BytesIO(source_bytes))
            for page_number, page in enumerate(reader.pages, start=1):
                try:
                    lines = _pypdf_layout_lines(page)
                    if not _append_records_from_text_lines(
                        stations,
                        page_data_records,
                        lines,
                        source_file,
                        page_number,
                        "pypdf",
                    ) and _page_type(lines) == MAIN_REPORT_PAGE_TYPE:
                        errors.append(f"{source_file} หน้า {page_number}: ไม่พบข้อมูลสถานี")
                except Exception as error:
                    errors.append(f"{source_file} หน้า {page_number}: {error}")
        except Exception as error:
            errors.append(f"เปิดไฟล์ {source_file} ไม่สำเร็จ: {error}")
        return stations, page_data_records, errors

    try:
        text_reader = PdfReader(BytesIO(source_bytes))
        with pdfplumber.open(BytesIO(source_bytes)) as pdf:
            for page_number, page in enumerate(pdf.pages, start=1):
                try:
                    lines = _pypdf_layout_lines(text_reader.pages[page_number - 1])
                    ocr_used = False
                    if parser_mode == "auto" and not lines:
                        try:
                            lines = _ocr_page_lines(source_bytes, page_number)
                            ocr_used = True
                        except RuntimeError:
                            # Keep the established text-PDF workflow unchanged when OCR
                            # is not installed locally or a blank supporting page has no text.
                            continue
                    if _page_type(lines) != MAIN_REPORT_PAGE_TYPE:
                        page_data_records.extend(
                            extract_page_data_records(
                                lines,
                                source_file,
                                page_number,
                                parser_method="OCR (Tesseract)" if ocr_used else "pypdf",
                            )
                        )
                        continue
                    station = (
                        parse_text_lines(lines, source_file, page_number, "OCR (Tesseract)")
                        if ocr_used
                        else parse_page(page, source_file, page_number)
                    )
                    if not ocr_used and parser_mode == "auto" and _should_try_text_fallback(station):
                        text_station = parse_pypdf_page(
                            text_reader.pages[page_number - 1], source_file, page_number
                        )
                        station = merge_station_pages(station, text_station)
                    if any(station.data.values()) or station.frequencies:
                        stations.append(station)
                        page_data_records.extend(
                            extract_page_data_records(
                                lines, source_file, page_number, station, station.parser_method
                            )
                        )
                    else:
                        errors.append(f"{source_file} หน้า {page_number}: ไม่พบข้อมูลสถานี")
                except Exception as error:  # Keep processing the remaining report pages.
                    errors.append(f"{source_file} หน้า {page_number}: {error}")
    except Exception as error:
        errors.append(f"เปิดไฟล์ {source_file} ไม่สำเร็จ: {error}")
    return stations, page_data_records, errors


def parse_pdf_bytes(
    source_file: str, source_bytes: bytes, parser_mode: str = "auto"
) -> tuple[list[StationPage], list[str]]:
    """Compatibility wrapper that returns only the main-report station pages."""
    stations, _, errors = parse_document_bytes(source_file, source_bytes, parser_mode)
    return stations, errors


def page_to_summary_record(station: StationPage) -> dict[str, str]:
    data = station.data
    return {
        "ผู้ประกอบการ": data["operator"],
        "เลขที่ใบอนุญาตตั้ง": data["license"],
        "ที่ตั้ง": data["location"],
        "ตำบล": data["subdistrict"],
        "อำเภอ": data["district"],
        "จังหวัด": data["province"],
        "รหัสไปรษณีย์": data["zipcode"],
        "Longitude": data["longitude"],
        "Latitude": data["latitude"],
        "ความถี่": join_frequency_values(station.frequencies, "frequency"),
        "ตราอักษร": join_frequency_values(station.frequencies, "brand"),
        "รุ่น/แบบ": join_frequency_values(station.frequencies, "model"),
        "กำลังส่ง (วัตต์)": join_frequency_values(station.frequencies, "power"),
        "อัตราขยายสายอากาศ (dBi)": join_frequency_values(station.frequencies, "gain"),
        "ความสูงสายอากาศ (เมตร)": join_frequency_values(station.frequencies, "height"),
        "ระยะห่างจากเสา ที่ตั้งสายอากาศ": data["max_dist_text"],
        "ระยะที่วัด/คำนวณ (เมตร)": data["max_dist_val"],
        "ระดับการแผ่คลื่นแม่เหล็กไฟฟ้าสูงสุด": data["max_rad"],
        "วันที่วัด/คำนวณ": data["date_calc"],
        "ลงชื่อ": data["signature"],
        "วันที่รายงาน": data["date_report"],
    }


def station_to_detail_records(station: StationPage) -> list[dict[str, str]]:
    base = page_to_summary_record(station)
    if not station.frequencies:
        return [base]

    records: list[dict[str, str]] = []
    for frequency in station.frequencies:
        record = base.copy()
        record.update(
            {
                "ความถี่": frequency["frequency"],
                "ตราอักษร": frequency["brand"],
                "รุ่น/แบบ": frequency["model"],
                "กำลังส่ง (วัตต์)": frequency["power"],
                "อัตราขยายสายอากาศ (dBi)": frequency["gain"],
                "ความสูงสายอากาศ (เมตร)": frequency["height"],
            }
        )
        records.append(record)
    return records


def build_report_records(stations: Sequence[StationPage], detail_mode: bool) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    for station in stations:
        page_records = station_to_detail_records(station) if detail_mode else [page_to_summary_record(station)]
        records.extend(page_records)

    for index, record in enumerate(records, start=1):
        record["ลำดับที่"] = index
    return records


def build_qa_records(stations: Sequence[StationPage]) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []
    for station in stations:
        status = "พร้อมใช้งาน" if not station.missing_fields else "ต้องตรวจสอบ"
        records.append(
            {
                "ไฟล์ต้นทาง": station.source_file,
                "หน้า": station.page_number,
                "วิธีอ่าน PDF": station.parser_method,
                "สถานะ": status,
                "ข้อมูลที่ไม่พบ": ", ".join(station.missing_fields),
                "จำนวนความถี่": len(station.frequencies),
                "หมายเหตุ": station.note,
            }
        )
    return records


def build_page_data_records(records: Sequence[PageDataRecord]) -> list[dict[str, object]]:
    """Build the row-based supporting-document output requested for each page type."""
    output: list[dict[str, object]] = []
    for record in records:
        missing_fields = [
            field for field in record.required_fields if not record.data.get(field, "")
        ]
        output.append(
            {
                "ไฟล์ต้นทาง": record.source_file,
                "หน้าต้นทาง": record.page_number,
                "หน้าที่พบ": record.page_type,
                **{
                    column: record.data.get(column, "")
                    for column in PAGE_DATA_COLUMNS
                    if column not in {"ไฟล์ต้นทาง", "หน้าต้นทาง", "หน้าที่พบ", "สถานะ", "ข้อมูลที่ไม่พบ"}
                },
                "สถานะ": "พร้อมใช้งาน" if not missing_fields else "ต้องตรวจสอบ",
                "ข้อมูลที่ไม่พบ": ", ".join(missing_fields),
            }
        )
    return output


def create_csv_bytes(records: Sequence[dict[str, object]], columns: Sequence[str]) -> bytes:
    """Create a UTF-8-with-BOM CSV so Thai text opens correctly in Excel."""
    output = StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(records)
    return output.getvalue().encode("utf-8-sig")


def _add_excel_sheet(workbook: Workbook, title: str, records: Sequence[dict[str, object]], columns: Sequence[str]) -> None:
    """Add a clean, filterable worksheet without changing the extracted values."""
    worksheet = workbook.create_sheet(title)
    worksheet.append(list(columns))
    for record in records:
        worksheet.append([record.get(column, "") for column in columns])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 34

    for column_index, cells in enumerate(worksheet.iter_cols(), start=1):
        longest_value = max((len(str(cell.value or "")) for cell in cells), default=10)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(12, min(longest_value + 2, 42))


def _add_license_check_columns(worksheet) -> None:
    """Add ready-to-use licence lookup columns beside the report data."""
    input_column = len(REPORT_COLUMNS) + 1
    result_column = input_column + 1
    input_letter = get_column_letter(input_column)
    result_letter = get_column_letter(result_column)
    report_last_row = worksheet.max_row
    check_last_row = max(report_last_row, 1001)
    license_range = f"$C$2:$C${max(report_last_row, 2)}"

    header_fill = PatternFill("solid", fgColor="C55A11")
    result_fill = PatternFill("solid", fgColor="E2F0D9")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    header_font = Font(color="FFFFFF", bold=True)

    input_header = worksheet.cell(1, input_column, "เลขที่ใบอนุญาตที่ต้องการตรวจสอบ\n(วางเลขที่นี่)")
    result_header = worksheet.cell(1, result_column, "ผลการตรวจสอบ")
    for cell in (input_header, result_header):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_number in range(2, check_last_row + 1):
        input_cell = worksheet.cell(row_number, input_column)
        input_cell.number_format = "@"
        input_cell.fill = input_fill
        input_cell.alignment = Alignment(vertical="top", wrap_text=True)

        result_cell = worksheet.cell(row_number, result_column)
        result_cell.value = (
            f'=IF({input_letter}{row_number}="","",'
            f'IF(COUNTIFS({license_range},{input_letter}{row_number})>0,"พบ","ไม่พบ"))'
        )
        result_cell.fill = result_fill
        result_cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    worksheet.column_dimensions[input_letter].width = 31
    worksheet.column_dimensions[result_letter].width = 18
    worksheet.row_dimensions[1].height = 42
    worksheet.auto_filter.ref = f"A1:V{report_last_row}"


def create_excel_bytes(
    report_records: Sequence[dict[str, object]],
    qa_records: Sequence[dict[str, object]],
    page_data_records: Sequence[dict[str, object]] = (),
) -> bytes:
    """Create an Excel report with main data, page-based data and review records."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _add_excel_sheet(workbook, "Report_Data", report_records, REPORT_COLUMNS)
    _add_license_check_columns(workbook["Report_Data"])
    _add_excel_sheet(workbook, "ข้อมูลตามหน้า", page_data_records, PAGE_DATA_COLUMNS)
    _add_excel_sheet(workbook, "ตรวจสอบข้อมูล", qa_records, QA_COLUMNS)

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
